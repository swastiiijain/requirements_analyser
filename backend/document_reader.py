"""Utilities for detecting the active document window and extracting its text contents."""
from __future__ import annotations

import os
from pathlib import Path
from typing import Optional, Tuple

import pygetwindow as gw
import psutil  # type: ignore

try:
    import win32gui, win32process, win32com.client  # type: ignore
except ImportError:
    # On non-Windows OS this module won't import, but the util is Windows-only anyway.
    pass

# Supported document extensions we know how to read
SUPPORTED_EXTENSIONS = {".docx", ".xlsx", ".pdf"}


# ---------------------------------------------------------------------------
# Helpers for path extraction
# ---------------------------------------------------------------------------


def _extract_from_title(title: str) -> Optional[str]:
    """Find a pathname in the window title (best-effort heuristic)."""
    lowered = title.lower()
    for ext in SUPPORTED_EXTENSIONS:
        if ext in lowered:
            parts = title.replace(" – ", " - ").split(" - ")
            for p in parts:
                if p.lower().strip().endswith(ext):
                    return p.strip()
    return None


def _extract_from_cmdline(pid: int) -> Optional[str]:
    """Inspect process command-line for a supported document path."""
    try:
        proc = psutil.Process(pid)
        for token in proc.cmdline():
            # Word wraps the path in quotes sometimes → strip them
            token = token.strip('"')
            if os.path.splitext(token)[1].lower() in SUPPORTED_EXTENSIONS and Path(token).exists():
                return token
    except Exception:
        pass
    return None


def _extract_from_office_com(process_name: str) -> Optional[str]:
    """Use COM Automation to query active document path for Word/Excel."""
    try:
        if process_name.lower() == "winword.exe":
            word = win32com.client.GetActiveObject("Word.Application")  # type: ignore
            doc = word.ActiveDocument  # pyright: ignore[reportAny]
            return doc.FullName if doc else None
        if process_name.lower() == "excel.exe":
            excel = win32com.client.GetActiveObject("Excel.Application")  # type: ignore
            wb = excel.ActiveWorkbook  # pyright: ignore[reportAny]
            return wb.FullName if wb else None
    except Exception:
        # COM call failed (e.g., no running instance or security restrictions)
        return None
    return None


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def get_active_document_text() -> Tuple[str, str]:
    """Return (absolute_path, extracted_text) for the current foreground document."""
    active = gw.getActiveWindow()
    if not active:
        raise FileNotFoundError("No active window detected.")

    title = active.title or ""

    # 1. Try COM or command-line inspection via PID
    hwnd = active._hWnd  # pygetwindow exposes private
    try:
        pid = win32process.GetWindowThreadProcessId(hwnd)[1]  # type: ignore
    except Exception:
        pid = None

    file_path: Optional[str] = None
    process_name = ""

    if pid:
        try:
            process_name = psutil.Process(pid).name()
        except Exception:
            process_name = ""

        # Office COM first (most reliable for Word/Excel)
        file_path = _extract_from_office_com(process_name)

        # Next: command-line args
        if not file_path:
            file_path = _extract_from_cmdline(pid)

    # 2. Heuristic window-title fallback
    if not file_path:
        file_path = _extract_from_title(title)

    if not file_path or not Path(file_path).exists():
        raise FileNotFoundError("Could not determine path of active document.")

    file_path = os.path.abspath(file_path)
    ext = os.path.splitext(file_path)[1].lower()

    text: str = ""
    if ext == ".docx":
        from docx import Document  # type: ignore

        doc = Document(file_path)
        text = "\n".join(p.text for p in doc.paragraphs)
    elif ext == ".xlsx":
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(file_path, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                line = " ".join(str(c) for c in row if c is not None)
                text += line + "\n"
    elif ext == ".pdf":
        import pdfplumber  # type: ignore

        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text() or ""
                text += extracted + "\n"
    else:
        raise FileNotFoundError(f"Unsupported extension: {ext}")

    return file_path, text 